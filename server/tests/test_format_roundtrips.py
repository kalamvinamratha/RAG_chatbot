"""Generated, synthetic parser fixtures (no real patient documents)."""

import csv

import docx
import openpyxl
import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches

from app.ingestion.chunkers.semantic_chunker import MedicalSemanticChunker
from app.ingestion.parsers.factory import get_parser_for_file


@pytest.mark.parametrize("extension", ["pdf", "docx", "xlsx", "csv", "pptx", "txt", "md"])
def test_supported_formats_preserve_synthetic_clinical_values(tmp_path, extension):
    path = tmp_path / f"synthetic.{extension}"
    note = "SYNTHETIC TEST ONLY. HbA1c 7.2 percent. Metformin 500 mg twice daily."
    if extension == "pdf":
        with pymupdf.open() as document:
            document.new_page().insert_text((50, 50), note)
            document.save(path)
    elif extension == "docx":
        document = docx.Document()
        document.add_paragraph(note)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Test"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "HbA1c"
        table.cell(1, 1).text = "7.2"
        document.save(path)
    elif extension == "xlsx":
        workbook = openpyxl.Workbook()
        workbook.active.append(["Synthetic note", "Value"])
        workbook.active.append([note, 7.2])
        workbook.save(path)
        workbook.close()
    elif extension == "csv":
        with path.open("w", newline="", encoding="utf-8-sig") as output:
            csv.writer(output).writerows([["Synthetic note", "Value"], [note, "7.2"]])
    elif extension == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = note
        presentation.save(path)
    else:
        path.write_text(note)
    parsed = get_parser_for_file(path).parse(path)
    chunks = MedicalSemanticChunker().chunk_document(parsed)
    content = "\n".join(c.content for c in chunks)
    assert parsed.total_pages == 1
    assert "7.2" in content
    assert "500" in content
    assert "Metformin" in content


def test_excel_zero_values_are_not_discarded(tmp_path):
    path = tmp_path / "zeros.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Synthetic event count", "Value"])
    workbook.active.append([0, 0])
    workbook.save(path)
    workbook.close()
    parsed = get_parser_for_file(path).parse(path)
    assert "| 0 | 0 |" in parsed.pages[0].tables[0]
