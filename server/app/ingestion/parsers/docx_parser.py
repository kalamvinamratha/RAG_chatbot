from pathlib import Path

import docx

from app.ingestion.models import ParsedDocument, ParsedPage
from app.ingestion.parsers.base_parser import BaseDocumentParser


class DocxDocumentParser(BaseDocumentParser):
    def parse(self, file_path: Path | str) -> ParsedDocument:
        path = Path(file_path)
        doc = docx.Document(str(path))

        text_parts: list[str] = []
        tables: list[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                text_parts.append(text)

        for t in doc.tables:
            rows_data: list[list[str]] = []
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells]
                rows_data.append(cells)

            if rows_data:
                headers = rows_data[0]
                md_table = (
                    "| "
                    + " | ".join(headers)
                    + " |\n| "
                    + " | ".join(["---"] * len(headers))
                    + " |\n"
                )
                for r in rows_data[1:]:
                    md_table += "| " + " | ".join(r) + " |\n"
                tables.append(md_table)

        full_text = "\n\n".join(text_parts)
        pages = [
            ParsedPage(
                page_number=1,
                text=full_text,
                tables=tables,
            )
        ]

        return ParsedDocument(
            filename=path.name,
            file_type="docx",
            total_pages=1,
            pages=pages,
            metadata={"file_size": path.stat().st_size},
        )
