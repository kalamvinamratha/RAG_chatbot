from pathlib import Path

import openpyxl

from app.ingestion.models import ParsedDocument, ParsedPage
from app.ingestion.parsers.base_parser import BaseDocumentParser


class ExcelDocumentParser(BaseDocumentParser):
    def parse(self, file_path: Path | str) -> ParsedDocument:
        path = Path(file_path)
        wb = openpyxl.load_workbook(str(path), data_only=True)
        pages: list[ParsedPage] = []

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            rows: list[list[str]] = []

            for row in ws.iter_rows(values_only=True):
                if any(cell is not None and cell != "" for cell in row):
                    clean_row = [str(cell) if cell is not None else "" for cell in row]
                    rows.append(clean_row)

            if not rows:
                continue

            headers = rows[0]
            md_table = (
                f"### Sheet: {sheet_name}\n\n"
                + "| "
                + " | ".join(headers)
                + " |\n| "
                + " | ".join(["---"] * len(headers))
                + " |\n"
            )
            for r in rows[1:]:
                # Pad row to match header length
                padded_r = r + [""] * (len(headers) - len(r))
                md_table += "| " + " | ".join(padded_r[: len(headers)]) + " |\n"

            pages.append(
                ParsedPage(
                    page_number=sheet_idx + 1,
                    text=f"Clinical Lab / Tabular Dataset for Sheet '{sheet_name}':\n\n{md_table}",
                    tables=[md_table],
                    metadata={"sheet_name": sheet_name},
                )
            )

        wb.close()

        return ParsedDocument(
            filename=path.name,
            file_type="xlsx",
            total_pages=len(pages),
            pages=pages,
            metadata={"file_size": path.stat().st_size},
        )
