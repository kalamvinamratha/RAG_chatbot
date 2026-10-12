from pathlib import Path

from pptx import Presentation

from app.ingestion.models import ParsedDocument, ParsedPage
from app.ingestion.parsers.base_parser import BaseDocumentParser


class PPTXDocumentParser(BaseDocumentParser):
    def parse(self, file_path: Path | str) -> ParsedDocument:
        path = Path(file_path)
        prs = Presentation(str(path))
        pages: list[ParsedPage] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_number = slide_idx + 1
            text_chunks: list[str] = []
            tables: list[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        text_chunks.append(text)
                elif shape.has_table:
                    table_rows: list[list[str]] = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells]
                        table_rows.append(row_cells)
                    if table_rows:
                        headers = table_rows[0]
                        md_table = (
                            "| "
                            + " | ".join(headers)
                            + " |\n| "
                            + " | ".join(["---"] * len(headers))
                            + " |\n"
                        )
                        for r in table_rows[1:]:
                            md_table += "| " + " | ".join(r) + " |\n"
                        tables.append(md_table)

            # Slide notes
            notes_text = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    text_chunks.append(f"Speaker Notes: {notes_text}")

            combined_text = f"Slide {slide_number}:\n" + "\n".join(text_chunks)
            pages.append(
                ParsedPage(
                    page_number=slide_number,
                    text=combined_text,
                    tables=tables,
                    metadata={"slide_number": slide_number},
                )
            )

        return ParsedDocument(
            filename=path.name,
            file_type="pptx",
            total_pages=len(pages),
            pages=pages,
            metadata={"file_size": path.stat().st_size},
        )
